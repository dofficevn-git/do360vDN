# -*- coding: utf-8 -*-
"""Sinh BIỂU MẪU EXCEL CHUẨN cho từng bộ phận từ danh mục chỉ tiêu (config/indicators/*.yaml) + luật bóc tách khớp sẵn.
   python tools/make_forms.py            → config/templates/forms/bieu-mau-<bo_phan>.xlsx + config/extract-rules/bieu_mau_<bo_phan>_xlsx.yaml
Xã ban hành biểu mẫu này cho bộ phận điền (cột "Giá trị"), nộp lại qua web hoặc chat → bóc đúng 100% không cần OCR/LLM
(khuyến nghị "phát biểu mẫu Excel chuẩn" trong tài liệu thiết kế). Chỉ tiêu capture_method = api (hệ thống tự cấp) không đưa vào mẫu.
Mỗi lần đổi danh mục: chạy lại; luật tăng version nếu nội dung đổi (file YAML có version — bất biến theo version)."""
import glob, io, os, re, sys
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(__file__), "..")
CFG = os.path.join(ROOT, "config")
DEPT_NAME = {
    "van_phong": "Văn phòng HĐND và UBND", "kinh_te": "Phòng Kinh tế, Hạ tầng và Đô thị", "van_hoa_xa_hoi": "Phòng Văn hóa - Xã hội",
    "hanh_chinh_cong": "Trung tâm phục vụ hành chính công", "cong_an": "Công an", "y_te": "Trạm Y tế", "quan_su": "Ban Chỉ huy quân sự",
}

def load_indicators():
    out = []
    for f in sorted(glob.glob(os.path.join(CFG, "indicators", "*.y*ml"))):
        doc = yaml.safe_load(io.open(f, encoding="utf-8")) or {}
        for it in doc.get("indicators") or []:
            if it.get("is_active", True) is False or (it.get("capture_method") or doc.get("capture_method")) == "api":
                continue
            out.append({"id": it["id"], "name": it["name"], "unit": it.get("unit", ""), "domain": doc.get("domain"),
                        "owner_dept": it.get("owner_dept") or doc.get("owner_dept") or "van_phong", "value_type": it.get("value_type", "integer")})
    return out

def make_form(dept, items, path):
    wb = Workbook(); ws = wb.active; ws.title = "Biểu mẫu"
    bold = Font(name="Times New Roman", bold=True, size=12); normal = Font(name="Times New Roman", size=12)
    thin = Side(style="thin", color="999999"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "UBND XÃ/PHƯỜNG: ……………………"; ws["A1"].font = bold
    ws["A2"] = f"BIỂU MẪU SỐ LIỆU ĐỊNH KỲ — {DEPT_NAME.get(dept, dept).upper()}"; ws["A2"].font = Font(name="Times New Roman", bold=True, size=14)
    ws["A3"] = "Kỳ báo cáo (vd 2026-Q3, 2026-M09, 2026-9M, 2026-Y):"; ws["D3"] = ""; ws["A3"].font = normal
    ws["A4"] = "Ngày chốt số liệu (dd/mm/yyyy):"; ws["D4"] = ""; ws["A4"].font = normal
    ws["A5"] = "Chỉ điền cột 'Giá trị' bằng số (không gõ đơn vị, không gõ chữ). Không sửa cột 'Mã chỉ tiêu'. Không dán danh sách có họ tên, CCCD, số điện thoại."
    ws["A5"].font = Font(name="Times New Roman", italic=True, size=11, color="7F7F7F")
    head = ["STT", "Mã chỉ tiêu", "Tên chỉ tiêu", "Đơn vị tính", "Giá trị", "Ghi chú"]
    ws.append([]); ws.append(head)
    hr = ws.max_row
    for c in range(1, 7):
        cell = ws.cell(row=hr, column=c); cell.font = bold; cell.border = border
        cell.fill = PatternFill("solid", fgColor="DDEBF7"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle="Chỉ nhập số", error="Ô này chỉ nhận số (dùng dấu phẩy cho phần thập phân).")
    ws.add_data_validation(dv)
    for i, it in enumerate(items, start=1):
        ws.append([i, it["id"], it["name"], it["unit"], None, ""])
        r = ws.max_row
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c); cell.font = normal; cell.border = border
        ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(row=r, column=5).number_format = "#,##0.##" if it["value_type"] in ("decimal", "percent") else "#,##0"
        dv.add(ws.cell(row=r, column=5))
    ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 16; ws.column_dimensions["F"].width = 30
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.protection.sheet = False
    wb.save(path)

def make_rule(dept, items, path):
    st = f"bieu_mau_{dept}_xlsx"
    body = {"source_type": st, "version": 1, "file_kind": "xlsx", "domain": None,
            "match": {"sheet": "Biểu mẫu", "header_contains": "Mã chỉ tiêu"},
            "map": [{"row": it["id"], "column": "Giá trị", "indicator": it["id"]} for it in items]}
    header = (f"# Luật bóc tách BIỂU MẪU EXCEL CHUẨN của bộ phận {DEPT_NAME.get(dept, dept)} — SINH TỰ ĐỘNG bởi tools/make_forms.py, không sửa tay.\n"
              f"# Khớp theo cột 'Mã chỉ tiêu' (nhãn hàng = mã) và cột 'Giá trị'. Đổi danh mục → chạy lại script và tăng version.\n")
    if os.path.exists(path):
        old = yaml.safe_load(io.open(path, encoding="utf-8")) or {}
        if old.get("map") != body["map"]:
            body["version"] = int(old.get("version", 1)) + 1
        else:
            body["version"] = int(old.get("version", 1))
    io.open(path, "w", encoding="utf-8", newline="\n").write(header + yaml.safe_dump(body, allow_unicode=True, sort_keys=False, width=200))

def main():
    items = load_indicators()
    by_dept = {}
    for it in items:
        by_dept.setdefault(it["owner_dept"], []).append(it)
    os.makedirs(os.path.join(CFG, "templates", "forms"), exist_ok=True)
    for dept, its in sorted(by_dept.items()):
        xlsx = os.path.join(CFG, "templates", "forms", f"bieu-mau-{dept.replace('_', '-')}.xlsx")
        rule = os.path.join(CFG, "extract-rules", f"bieu_mau_{dept}_xlsx.yaml")
        make_form(dept, its, xlsx); make_rule(dept, its, rule)
        print(f"{dept}: {len(its)} chỉ tiêu → {os.path.relpath(xlsx, ROOT)} + {os.path.relpath(rule, ROOT)}")

if __name__ == "__main__":
    main()
