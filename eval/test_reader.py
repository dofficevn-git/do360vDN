# -*- coding: utf-8 -*-
"""Bộ đánh giá bóc tách (chạy không cần server):  python eval/test_reader.py
Tự dựng file mẫu theo bố cục trong luật (khi chưa có file thật của xã), rồi kiểm:
- chuẩn hoá số VN, phát hiện mơ hồ
- Excel ánh xạ ô + từ điển nhãn; Word mẫu câu; PDF có lớp chữ
- luật không khớp → không bịa số
Khi có file thật: bỏ vào eval/private/<source_type>/*.xlsx kèm expected.json, script sẽ chấm thêm."""
import io, json, os, sys, glob
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "services", "reader"))
from app.numbers import parse_vn_number
from app.readers import read_any
from app.rules import apply_rule
import yaml

FAILS = []
def check(name, cond, detail=""):
    print(("  OK   " if cond else "  FAIL ") + name + (f" — {detail}" if detail and not cond else ""))
    if not cond: FAILS.append(name)

CFG = os.path.join(os.path.dirname(__file__), "..", "config")
RULE = yaml.safe_load(open(os.path.join(CFG, "extract-rules", "mot_cua_tong_hop_xlsx.yaml"), encoding="utf-8"))
TYPES = {i["id"]: i.get("value_type", "integer") for f in glob.glob(os.path.join(CFG, "indicators", "*.yaml")) for i in yaml.safe_load(open(f, encoding="utf-8"))["indicators"]}
NUMS = {"tthc.tiep_nhan": 17206, "tthc.giai_quyet": 17007, "tthc.truoc_han": 16979, "tthc.qua_han": 28, "tthc.dang_xu_ly": 199}

print("1. Chuẩn hoá số kiểu Việt Nam")
for raw, exp, amb in [("17.206", 17206, False), ("17206", 17206, False), ("253,515 tỷ đồng", 253_515_000_000, False), ("96,3%", 96.3, False),
                      ("1.234.567", 1234567, False), ("352.732 tỷ đồng", 352_732_000_000_000, True), ("12,5", 12.5, True), ("8,2 triệu đồng", 8_200_000, False)]:
    p = parse_vn_number(raw, "integer")
    check(f"'{raw}' → {exp}{' (mơ hồ)' if amb else ''}", abs(p.value - exp) < 1e-6 and p.ambiguous == amb, f"{p}")

print("2. Excel theo biểu mẫu một cửa (ánh xạ ô + từ điển nhãn)")
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tổng hợp"
ws["A1"] = "BÁO CÁO TỔNG HỢP TIẾP NHẬN VÀ GIẢI QUYẾT HỒ SƠ"; ws["B2"] = "15/09/2026"
ws.append([]); ws.append(["STT", "Chỉ tiêu", "Trong kỳ", "Lũy kế"])
rows = [("1", "Tổng số hồ sơ tiếp nhận", "17.206", "17.206"), ("2", "Đã trả kết quả", "17.007", "17.007"),
        ("2.1", "Đã giải quyết trước hạn", 16979, 16979), ("2.2", "Đã giải quyết quá hạn", "28", "28"), ("3", "Hồ sơ đang xử lý", "199", "199")]
for r in rows: ws.append(list(r))
buf = io.BytesIO(); wb.save(buf)
doc = read_any("tong-hop.xlsx", buf.getvalue())
res = apply_rule(doc, RULE, TYPES)
got = {e.indicator_id: e.value for e in res.records}
check("khớp luật, bóc đủ 5/5 chỉ tiêu đúng giá trị", res.matched and got == NUMS, f"{got} misses={res.misses}")
check("source_ref trỏ đúng ô (Trong kỳ = cột C)", all(e.source_ref["cell"].startswith("C") for e in res.records), str([e.source_ref["cell"] for e in res.records]))
check("'Đã trả kết quả' được từ điển nhãn ánh xạ về 'Đã giải quyết'", "tthc.giai_quyet" in got)

print("3. Excel không khớp luật → không bịa")
wb2 = openpyxl.Workbook(); ws2 = wb2.active; ws2.title = "Sheet1"; ws2.append(["Tên", "Giá trị"]); ws2.append(["Số hộ nghèo", 12])
b2 = io.BytesIO(); wb2.save(b2)
res2 = apply_rule(read_any("khac.xlsx", b2.getvalue()), RULE, TYPES)
check("không khớp → matched=False, 0 bản ghi", not res2.matched and not res2.records, f"{res2}")

print("4. Word — mẫu câu")
from docx import Document
d = Document(); d.add_paragraph("Trong 9 tháng đầu năm 2026, Bộ phận một cửa đã tiếp nhận 17.206 hồ sơ; giải quyết 17.007 hồ sơ, trong đó 16.979 hồ sơ trước hạn, 28 hồ sơ quá hạn; 199 hồ sơ đang trong thời hạn giải quyết.")
d.add_paragraph("Thu ngân sách đạt 352.732 tỷ đồng/304,252 tỷ đồng dự toán.")
bd = io.BytesIO(); d.save(bd)
RULE_DOCX = {"source_type": "bao_cao_mot_cua_docx", "version": 1, "file_kind": "docx", "match": {"text_contains": "một cửa"},
             "sentences": [{"pattern": "tiếp nhận {a:number} hồ sơ; giải quyết {b:number} hồ sơ, trong đó {c:number} hồ sơ trước hạn, {d:number} hồ sơ quá hạn; {e:number} hồ sơ đang",
                            "groups": {"a": "tthc.tiep_nhan", "b": "tthc.giai_quyet", "c": "tthc.truoc_han", "d": "tthc.qua_han", "e": "tthc.dang_xu_ly"}}]}
res3 = apply_rule(read_any("bc.docx", bd.getvalue()), RULE_DOCX, TYPES)
got3 = {e.indicator_id: e.value for e in res3.records}
check("bóc 5 số từ câu văn", got3 == NUMS, f"{got3} misses={res3.misses}")
p = parse_vn_number("352.732 tỷ đồng")
check("'352.732 tỷ đồng' bị đánh dấu mơ hồ (không tự quyết)", p.ambiguous)

print("5. PDF có lớp chữ")
import pymupdf
pdf = pymupdf.open(); page = pdf.new_page()
page.insert_text((50, 80), "Bo phan mot cua da tiep nhan 17.206 ho so; giai quyet 17.007 ho so, trong do 16.979 ho so truoc han, 28 ho so qua han; 199 ho so dang trong thoi han", fontsize=9)
bp = pdf.tobytes()
docp = read_any("bc.pdf", bp, allow_ocr=False)
check("PDF nhận là pdf_text (có lớp chữ)", docp.file_kind == "pdf_text", docp.file_kind)
RULE_PDF = dict(RULE_DOCX, source_type="bao_cao_mot_cua_pdf", file_kind="pdf_text", match={"text_contains": "mot cua"},
                sentences=[{"pattern": "tiep nhan {a:number} ho so; giai quyet {b:number} ho so, trong do {c:number} ho so truoc han, {d:number} ho so qua han; {e:number} ho so dang", "groups": RULE_DOCX["sentences"][0]["groups"]}])
res5 = apply_rule(docp, RULE_PDF, TYPES)
check("bóc 5 số từ PDF, có bbox", {e.indicator_id: e.value for e in res5.records} == NUMS and all(e.source_ref.get("bbox") for e in res5.records), f"{[(e.indicator_id, e.value) for e in res5.records]} {res5.misses}")

print("6. File thật trong eval/private (nếu có)")
priv = os.path.join(os.path.dirname(__file__), "private")
n = 0
for exp_path in glob.glob(os.path.join(priv, "*", "expected.json")):
    exp = json.load(open(exp_path, encoding="utf-8"))
    rule = yaml.safe_load(open(os.path.join(CFG, "extract-rules", exp["rule"] + ".yaml"), encoding="utf-8"))
    for fpath in glob.glob(os.path.join(os.path.dirname(exp_path), "*.*")):
        if fpath.endswith("expected.json"): continue
        r = apply_rule(read_any(os.path.basename(fpath), open(fpath, "rb").read()), rule, TYPES)
        got = {e.indicator_id: e.value for e in r.records}
        ok = sum(1 for k, v in exp["values"].items() if got.get(k) == v)
        check(f"{os.path.basename(fpath)}: {ok}/{len(exp['values'])} đúng", ok == len(exp["values"]), str(got)); n += 1
if n == 0: print("  (chưa có file thật — bỏ qua)")

print(); print("KẾT QUẢ:", "ĐẠT" if not FAILS else f"CÒN {len(FAILS)} LỖI: {FAILS}")
sys.exit(1 if FAILS else 0)
