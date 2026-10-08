# -*- coding: utf-8 -*-
"""Kiểm chứng T1/T4 kênh chat, D5 QLVB (giả lập), D6 KPI, dashboard lãnh đạo.
   python eval/smoke_d5.py http://localhost:3070/api admin@do360.local <mật khẩu>
Dùng kỳ 2026-M08 (tháng) — độc lập với smoke_d1 (2026-9M) và smoke_d3 (2026-Q3).
Kho chỉ thêm không xoá → các bộ smoke chạy trên CSDL mới (infra: docker compose down -v; up -d); chạy lần hai sẽ báo trùng số."""
import base64, io, json, sys, urllib.request, uuid, datetime as dt

API, EMAIL, PW = sys.argv[1], sys.argv[2], sys.argv[3]
FAILS = []

def call(method, path, body=None, token=None):
    data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
    req = urllib.request.Request(API + path, method=method, data=data)
    req.add_header("Content-Type", "application/json; charset=utf-8")
    if token: req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            t = r.read(); return r.status, (json.loads(t) if t else None)
    except urllib.error.HTTPError as e:
        t = e.read().decode("utf-8", "replace")
        try: return e.code, json.loads(t)
        except Exception: return e.code, t

def check(name, cond, detail=""):
    print(("  OK   " if cond else "  FAIL ") + name + (f" — {detail}" if detail and not cond else ""))
    if not cond: FAILS.append(name)

st, d = call("POST", "/admin/auth/login", {"email": EMAIL, "password": PW}); tok = d["access_token"]
st, me = call("GET", "/admin/auth/me", token=tok) if False else (0, None)
call("POST", "/admin/catalog/reload", token=tok)
PERIOD = "2026-M08"; CUTOFF = "2026-08-31"
NUMS = {"tthc.tiep_nhan": 5000, "tthc.giai_quyet": 4900, "tthc.truoc_han": 4880, "tthc.qua_han": 20, "tthc.dang_xu_ly": 100}

# dọn kỳ cũ của lần chạy trước (huỷ) để tạo lại
st, runs = call("GET", "/admin/report-runs", token=tok)
for r in runs:
    if r["period"] == PERIOD and r["status"] != "cancelled":
        call("POST", f"/admin/report-runs/{r['run_id']}/cancel", {"reason": "chạy lại kiểm thử smoke_d5"}, token=tok)

print("1. Kênh chat (T1/T4) — giả lập tin cán bộ gửi vào bot")
st, cs = call("GET", "/admin/chat/status", token=tok)
check("trạng thái kênh: zalo, có web_url", st == 200 and cs["channel"] == "zalo" and cs["web_url"].startswith("http"), str(cs))
st, t = call("POST", "/admin/chat/test", token=tok)
check("thẻ thử được ghi nhật ký (sent/failed/skipped)", st in (200, 201) and t["status"] in ("sent", "failed", "skipped"), str(t))
# nộp file qua chat kèm mã kỳ
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tổng hợp"
ws.append(["BÁO CÁO TỔNG HỢP MỘT CỬA " + uuid.uuid4().hex[:6]]); ws.append([]); ws.append(["STT", "Chỉ tiêu", "Trong kỳ", "Lũy kế"])
for r in [("1", "Tổng số hồ sơ tiếp nhận", "5.000"), ("2", "Đã trả kết quả", "4.900"), ("2.1", "Đã giải quyết trước hạn", 4880), ("2.2", "Đã giải quyết quá hạn", "25"), ("3", "Hồ sơ đang xử lý", "100")]:
    ws.append(list(r))
buf = io.BytesIO(); wb.save(buf)
st, inb = call("POST", "/admin/chat/inbound", {"type": "file", "text": f"Gửi số một cửa kỳ {PERIOD}", "file": {"name": "mot-cua-thang-8.xlsx", "base64": base64.b64encode(buf.getvalue()).decode(), "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}}, token=tok)
check("file qua chat → Bronze → bóc 5 số, bot trả lời", st in (200, 201) and inb["handled"] and inb.get("result", {}).get("status") == "extracted" and len(inb["result"]["records"]) == 5, str(inb)[:400])
recs = {r["indicator_id"]: r for r in inb["result"]["records"]}
st, cur = call("GET", f"/admin/facts?period={PERIOD}", token=tok)
check("5 số một cửa vào kho từ file export (chat)", all(any(r["indicator_id"] == k and r["capture_method"] == "export" for r in cur) for k in NUMS), str([(r["indicator_id"], r["status"]) for r in cur]))
# lệnh XN
st, xn = call("POST", "/admin/chat/inbound", {"type": "text", "text": f"XN {recs['tthc.tiep_nhan']['fact_id']}"}, token=tok)
check("lệnh 'XN <mã>' xác nhận số", st in (200, 201) and xn["handled"] and xn["result"]["status"] == "confirmed", str(xn)[:300])
for k in ("tthc.giai_quyet", "tthc.truoc_han", "tthc.dang_xu_ly"):
    call("POST", "/admin/chat/inbound", {"type": "text", "text": f"XN {recs[k]['fact_id']}"}, token=tok)
st, q = call("POST", "/admin/chat/inbound", {"type": "text", "text": f"XN {recs['tthc.qua_han']['fact_id']}"}, token=tok)
check("quá hạn=25 sai số học → bot báo CÁCH LY", q["handled"] and "CÁCH LY" in (q["reply"] or ""), str(q)[:300])
st, msgs = call("GET", "/admin/chat/messages?limit=20", token=tok)
check("thẻ cách ly gửi Data Owner được ghi nhật ký (reference_type fact)", any(m["direction"] == "out" and m["reference_type"] == "fact" for m in msgs), str([(m["direction"], m["kind"], m["reference_type"]) for m in msgs][:8]))
st, su = call("POST", "/admin/chat/inbound", {"type": "text", "text": f"SUA {recs['tthc.qua_han']['fact_id']} = 20 bóc nhầm"}, token=tok)
check("lệnh 'SUA <mã> = 20' tạo phiên bản 2 confirmed", su["handled"] and su["result"]["version"] == 2 and su["result"]["status"] == "confirmed", str(su)[:300])
st, bad = call("POST", "/admin/chat/inbound", {"type": "text", "text": "xin chào"}, token=tok)
check("tin thường không phải lệnh → handled=false (đi tiếp luồng QR)", bad["handled"] is False, str(bad))
st, hp = call("POST", "/admin/chat/inbound", {"type": "text", "text": "HELP"}, token=tok)
check("HELP trả hướng dẫn lệnh", hp["handled"] and "XN" in hp["reply"], str(hp)[:200])
# thẻ form: nhập số không có file (lệnh NHAP) — ghi vào kỳ khác để không đụng bộ số tháng 8
st, nh = call("POST", "/admin/chat/inbound", {"type": "text", "text": "NHAP tthc.tiep_nhan = 1.234 2026-M07 31/07/2026"}, token=tok)
check("lệnh 'NHAP <chỉ tiêu> = <số> <kỳ> <ngày>' tạo số extracted nguồn form", st in (200, 201) and nh["handled"] and nh["result"]["status"] == "extracted", str(nh)[:300])
st, cur7 = call("GET", "/admin/facts?period=2026-M07", token=tok)
check("số NHAP có capture_method form, as_of_date 31/07", any(r["indicator_id"] == "tthc.tiep_nhan" and r["capture_method"] == "form" and r["as_of_date"] == "2026-07-31" and r["value"] == 1234 for r in cur7), str(cur7)[:300])
st, nh2 = call("POST", "/admin/chat/inbound", {"type": "text", "text": "NHAP tthc.khong_co = 5 2026-M07"}, token=tok)
check("NHAP chỉ tiêu không có trong danh mục → bot báo lỗi, không bịa", nh2["handled"] and "Không thực hiện được" in (nh2["reply"] or ""), str(nh2)[:200])
# dữ liệu cá nhân trong file nguồn → cảnh báo
wb2 = openpyxl.Workbook(); ws2 = wb2.active; ws2.title = "Tổng hợp"
ws2.append(["DANH SÁCH " + uuid.uuid4().hex[:6]]); ws2.append([]); ws2.append(["STT", "Chỉ tiêu", "Trong kỳ", "Ghi chú"])
ws2.append(["1", "Tổng số hồ sơ tiếp nhận", 10, "Nguyễn Văn A - CCCD 048090001234 - ĐT 0905123456"])
b2 = io.BytesIO(); wb2.save(b2)
st, pf = call("POST", "/admin/chat/inbound", {"type": "file", "text": "kỳ 2026-M07", "file": {"name": "danh-sach.xlsx", "base64": base64.b64encode(b2.getvalue()).decode()}}, token=tok)
check("file có CCCD/SĐT → cảnh báo dữ liệu cá nhân trong trả lời", pf["handled"] and "dữ liệu cá nhân" in (pf["reply"] or ""), str(pf)[:300])
st, fl = call("GET", "/admin/files?period=2026-M07", token=tok)
check("source_file.pii_warning được ghi", any(f.get("pii_warning") for f in fl), str([(f["original_name"], f.get("pii_warning")) for f in fl])[:300])
# Mini App: Zalo openMediaPicker POST file lên /public/do360/upload?token=…&period=… (không header Authorization)
wb3 = openpyxl.Workbook(); ws3 = wb3.active; ws3.title = "Tổng hợp"
ws3.append(["BÁO CÁO TỔNG HỢP MỘT CỬA " + uuid.uuid4().hex[:6]]); ws3.append([]); ws3.append(["STT", "Chỉ tiêu", "Trong kỳ", "Lũy kế"])
for r in [("1", "Tổng số hồ sơ tiếp nhận", 3000), ("2", "Đã trả kết quả", 2900), ("2.1", "Đã giải quyết trước hạn", 2890), ("2.2", "Đã giải quyết quá hạn", 10), ("3", "Hồ sơ đang xử lý", 100)]:
    ws3.append(list(r))
b3 = io.BytesIO(); wb3.save(b3)
bnd = "----zalo" + uuid.uuid4().hex
body3 = (f"--{bnd}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"mot-cua-thang-6.xlsx\"\r\nContent-Type: application/octet-stream\r\n\r\n").encode() + b3.getvalue() + f"\r\n--{bnd}--\r\n".encode()
req3 = urllib.request.Request(API + f"/public/do360/upload?token={tok}&period=2026-M06", method="POST", data=body3); req3.add_header("Content-Type", f"multipart/form-data; boundary={bnd}")
try:
    with urllib.request.urlopen(req3, timeout=180) as r: st, mu = r.status, json.loads(r.read())
except urllib.error.HTTPError as e: st, mu = e.code, e.read().decode("utf-8", "replace")
check("upload từ Mini App (token theo query) → Bronze channel miniapp → bóc 5 số", st == 200 and mu["extract"]["status"] == "extracted" and len(mu["extract"]["records"]) == 5, str(mu)[:300])
req4 = urllib.request.Request(API + "/public/do360/upload?token=sai&period=2026-M06", method="POST", data=body3); req4.add_header("Content-Type", f"multipart/form-data; boundary={bnd}")
try:
    urllib.request.urlopen(req4, timeout=60); st4 = 200
except urllib.error.HTTPError as e: st4 = e.code
check("upload với token sai → 401", st4 == 401, str(st4))

print("2. Mở kỳ, nhắc phòng còn thiếu, chốt, sinh bản nháp")
st, specs = call("GET", "/admin/catalog/report-specs", token=tok); spec_id = next(s["id"] for s in specs if s["spec_code"] == "ktxh_mot_cua")
st, run = call("POST", "/admin/report-runs", {"spec_id": spec_id, "period": PERIOD, "cutoff_date": CUTOFF}, token=tok)
check("mở kỳ 2026-M08", st in (200, 201) and run["status"] == "collecting", str(run)[:200]); rid = run["run_id"]
st, rm = call("POST", f"/admin/chat/remind/{rid}", token=tok)
check("nhắc: phòng hanh_chinh_cong còn chỉ tiêu chưa duyệt", st in (200, 201) and "hanh_chinh_cong" in rm["depts"] and len(rm["missing"]["hanh_chinh_cong"]) == 5, str(rm)[:300])
st, td = call("POST", "/admin/chat/inbound", {"type": "text", "text": "TIEN DO"}, token=tok)
check("lệnh TIEN DO trả tiến độ kỳ", td["handled"] and PERIOD in td["reply"], str(td)[:200])
st, cur = call("GET", f"/admin/facts?period={PERIOD}", token=tok)
call("POST", "/admin/facts/approve", {"ids": [r["id"] for r in cur if r["status"] == "confirmed"]}, token=tok)
for k in NUMS: call("POST", f"/admin/report-runs/{rid}/exceptions", {"indicator_id": k, "reason": "Một cửa chốt theo ngày xuất file"}, token=tok)
st, locked = call("POST", f"/admin/report-runs/{rid}/lock", token=tok)
check("chốt kỳ", st in (200, 201) and locked["status"] == "locked", str(locked)[:300])
st, rd = call("POST", f"/admin/report-runs/{rid}/render", token=tok)
check("sinh bản nháp", st in (200, 201) and rd["run"]["status"] == "drafted", str(rd)[:200])
# D4: tải lại bản Word đã biên tập để đối chiếu
import zipfile
def upload_docx(path, content, name):
    b = "----do360" + uuid.uuid4().hex
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\nContent-Type: application/octet-stream\r\n\r\n").encode() + content + f"\r\n--{b}--\r\n".encode()
    req = urllib.request.Request(API + path, method="POST", data=body); req.add_header("Content-Type", f"multipart/form-data; boundary={b}"); req.add_header("Authorization", "Bearer " + tok)
    try:
        with urllib.request.urlopen(req, timeout=180) as r: return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e: return e.code, e.read().decode("utf-8", "replace")
req = urllib.request.Request(API + f"/admin/report-runs/{rid}/download"); req.add_header("Authorization", "Bearer " + tok)
docx = urllib.request.urlopen(req).read()
st, ok = upload_docx(f"/admin/report-runs/{rid}/edited", docx, "ban-nhap-giu-nguyen.docx")
check("bản nháp tải lại nguyên vẹn → khớp snapshot, không số lạ", st in (200, 201) and ok["ok"] and not ok["unknown"], str(ok)[:300])
def tweak(docx_bytes, a, b):
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes)); out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == "word/document.xml": data = data.decode("utf-8").replace(a, b, 1).encode("utf-8")
        zout.writestr(it, data)
    zout.close(); return out.getvalue()
st, bad = upload_docx(f"/admin/report-runs/{rid}/edited", tweak(docx, "4.880", "4.808"), "ban-sua-sai-so.docx")
check("bản biên tập đổi 4.880 → 4.808: báo THIẾU 4.880 và số lạ 4.808", st in (200, 201) and not bad["ok"] and "4.880" in bad["missing"] and "4.808" in bad["unknown"], str(bad)[:300])
st, rr = call("GET", f"/admin/report-runs/{rid}", token=tok)
check("kết quả đối chiếu lưu trên run, có bản biên tập để tải", rr["run"]["edited_check"]["ok"] is False and rr["run"]["edited_file"], str(rr["run"].get("edited_check"))[:200])
st, ok2 = upload_docx(f"/admin/report-runs/{rid}/edited", tweak(docx, "hồ sơ đang trong thời hạn", "hồ sơ đang trong thời hạn (mục II.3.4)"), "ban-sua-loi-van.docx")
check("Văn phòng sửa lại đúng số (chỉ thêm lời văn, số mục II.3.4) → khớp, số mục không bị coi là số lạ", st in (200, 201) and ok2["ok"] and not ok2["unknown"], str(ok2)[:300])

print("2b. Đọc báo cáo (A2): cây tài liệu có mảnh số mang nguồn; lời dẫn có phiên bản + hậu kiểm")
st, pv = call("GET", f"/admin/report-runs/{rid}/preview", token=tok)
secs = pv["document"]["sections"]
check("preview: có mục II.3.4 với tên, câu số dạng mảnh", st == 200 and any(s["code"] == "II.3.4" and "Thủ tục hành chính" in s["title"] for s in secs), str([(s["code"], s["title"]) for s in secs])[:300])
slot = next(sl for s in secs for sl in s["slots"] if sl["slot_code"] == "tthc.mot_cua.tong_hop")
nums = [g for g in slot["segments"] if g["kind"] == "number"]
check("mọi mảnh số đều có nguồn (fact_id, ô nguồn, ngày chốt)", len(nums) == 5 and all(g["ref"].get("fact_id") and g["ref"].get("as_of_date") for g in nums), str(nums)[:300])
check("văn bản đọc được: text có 5.000 và 4.900", "5.000" in slot["text"] and "4.900" in slot["text"], slot["text"])
st, nl = call("GET", f"/admin/report-runs/{rid}/narratives", token=tok)
check("danh sách lời dẫn (mẫu rút gọn không có slot lời dẫn → danh sách rỗng, không lỗi)", st == 200 and isinstance(nl["slots"], list), str(nl)[:200])
# dùng mẫu đầy đủ để thử lời dẫn: mở kỳ 2026-M07 với ktxh_dinh_ky (không chốt, chỉ kiểm hậu kiểm lời dẫn)
spec_full = next(s["id"] for s in specs if s["spec_code"] == "ktxh_dinh_ky")
st, run7 = call("POST", "/admin/report-runs", {"spec_id": spec_full, "period": "2026-M07", "cutoff_date": "2026-07-31"}, token=tok)
if st == 400: st, rr7 = call("GET", "/admin/report-runs", token=tok); run7 = next(r for r in rr7 if r["period"] == "2026-M07" and r["status"] != "cancelled")
rid7 = run7["run_id"]
st, nl7 = call("GET", f"/admin/report-runs/{rid7}/narratives", token=tok)
check("mẫu đầy đủ có ≥ 4 slot lời dẫn kèm gợi ý khuôn câu", st == 200 and len(nl7["slots"]) >= 4 and any(s["suggestions"]["template"] for s in nl7["slots"]), str([(s["slot_code"], bool(s["suggestions"]["template"])) for s in nl7["slots"]])[:300])
st, bad = call("POST", f"/admin/report-runs/{rid7}/narratives/chi_dao.loi_dan", {"text": "Trong tháng 7 năm 2026 phường đã giải quyết 9.999 hồ sơ."}, token=tok)
check("lời dẫn có số lạ (9.999) → bị chặn 400", st == 400 and "số" in json.dumps(bad, ensure_ascii=False), str(bad)[:200])
st, warn = call("POST", f"/admin/report-runs/{rid7}/narratives/chi_dao.loi_dan", {"text": "Trong 6 tháng đầu năm 2026, UBND phường tập trung chỉ đạo các nhiệm vụ trọng tâm."}, token=tok)
check("lời dẫn ghi '6 tháng đầu năm' trong báo cáo tháng 7 → lưu được, có cảnh báo nhất quán kỳ", st in (200, 201) and warn["status"] == "draft" and warn["postcheck"]["warnings"], str(warn)[:300])
st, st7 = call("GET", f"/admin/report-runs/{rid7}", token=tok)
check("gate: cảnh báo nhóm nhất quán kỳ chỉ sau khi chấp nhận (bản nháp chưa tính)", not any(w["group"] == "nhat_quan" for w in st7["gate"]["warnings"]), str([w["group"] for w in st7["gate"]["warnings"]]))
st, acc = call("POST", f"/admin/report-runs/{rid7}/narratives/chi_dao.loi_dan/accept", token=tok)
check("chấp nhận → phiên bản mới status accepted", st in (200, 201) and acc["status"] == "accepted" and acc["version"] == 2, str(acc)[:200])
st, st7 = call("GET", f"/admin/report-runs/{rid7}", token=tok)
check("gate: có cảnh báo nhóm nhất quán kỳ sau khi chấp nhận", any(w["group"] == "nhat_quan" and w["kind"] == "period_word" for w in st7["gate"]["warnings"]), str(st7["gate"]["warnings"])[:300])
st, sg = call("POST", f"/admin/report-runs/{rid7}/narratives/kinh_te.loi_dan/suggest", {"provider": "template"}, token=tok)
check("gợi ý tự động (khuôn câu) → bản nháp nguồn template", st in (200, 201) and sg["source"] == "template" and sg["text"], str(sg)[:200])
st, g7 = call("GET", f"/admin/report-runs/{rid7}", token=tok)
check("gate nói bằng câu: lỗi 'Chưa có số' có tên chỉ tiêu, không có mã", any(e["kind"] == "missing" and "Chưa có số" in e["title"] and "tthc." not in e["title"] for e in g7["gate"]["errors"]), str(g7["gate"]["errors"][:2])[:300])
call("POST", f"/admin/report-runs/{rid7}/cancel", {"reason": "kết thúc kiểm thử lời dẫn"}, token=tok)

print("3. QLVB (D5, giả lập): sổ văn bản → sự kiện + giao việc; trình ký → ký → kỳ signed")
st, qs = call("GET", "/admin/qlvb/status", token=tok)
check("chế độ sim, configured", st == 200 and qs["mode"] == "sim" and qs["configured"], str(qs))
st, seed = call("POST", "/admin/qlvb/sim/documents", {"documents": [
    {"direction": "den", "so_ky_hieu": "1234/UBND-VP", "ngay": "2026-08-05", "loai": "Công văn", "trich_yeu": "V/v rà soát thủ tục hành chính một cửa", "nguoi_xu_ly_email": EMAIL, "han_xu_ly": "2026-08-15", "ngay_hoan_thanh": "2026-08-12"},
    {"direction": "den", "so_ky_hieu": "1250/STC-NS", "ngay": "2026-08-10", "loai": "Công văn", "trich_yeu": "V/v hướng dẫn quyết toán ngân sách", "nguoi_xu_ly_email": EMAIL, "han_xu_ly": "2026-08-20"},
    {"direction": "di", "so_ky_hieu": "88/KH-UBND", "ngay": "2026-08-18", "loai": "Kế hoạch", "trich_yeu": "Kế hoạch tiêm chủng mở rộng", "nguoi_ky": "Chủ tịch UBND"},
]}, token=tok)
check("nạp 3 văn bản giả lập", st in (200, 201) and seed["seeded"] == 3, str(seed))
st, sy = call("POST", "/admin/qlvb/sync", {"period": PERIOD}, token=tok)
check("đồng bộ: 3 văn bản → 3 sự kiện, 2 việc (1 đã hoàn thành)", st in (200, 201) and sy["fetched"] == 3 and sy["events_inserted"] == 3 and sy["tasks_created"] + sy["tasks_completed"] == 2, str(sy))
st, sy2 = call("POST", "/admin/qlvb/sync", {"period": PERIOD}, token=tok)
check("đồng bộ lại không nhân đôi sự kiện", sy2["events_inserted"] == 0 and sy2["events_skipped"] == 3, str(sy2))
st, lct = call("POST", "/admin/qlvb/sim/schedule", {"items": [{"date": "2026-08-20", "time": "08:00", "title": "Họp giao ban kinh tế - xã hội tháng 8", "location": "Hội trường UBND"}, {"date": "2026-08-25", "title": "Làm việc với Sở Tài chính về quyết toán ngân sách"}]}, token=tok)
check("lịch công tác → 2 sự kiện họp (MTG)", st in (200, 201) and lct["inserted"] == 2, str(lct))
st, ev8 = call("GET", f"/admin/events?period={PERIOD}", token=tok)
check("sự kiện họp có lĩnh vực theo luật (tai_chinh) và không nhân đôi khi nạp lại", any(e["object_type"] == "MTG" and e["domain"] == "tai_chinh" for e in ev8) and call("POST", "/admin/qlvb/sim/schedule", {"items": [{"date": "2026-08-20", "time": "08:00", "title": "Họp giao ban kinh tế - xã hội tháng 8"}]}, token=tok)[1]["skipped"] == 1, str([(e["object_type"], e["object_id"][:30], e["domain"]) for e in ev8])[:300])
# D6: nhắc việc sắp đến hạn (giao 1 việc hạn ngày mai)
st, tk = call("POST", "/admin/tasks", {"title": "Nộp số liệu tháng 9 cho Văn phòng", "assignee_id": 1, "due_at": (dt.datetime.now() + dt.timedelta(hours=30)).isoformat(), "priority": "high"}, token=tok)
st, rt = call("POST", "/admin/chat/remind-tasks", token=tok)
check("nhắc việc sắp đến hạn: ≥1 việc trong 48 giờ, 1 cán bộ", st in (200, 201) and rt["tasks"] >= 1 and rt["recipients"] >= 1, str(rt))
st, ev = call("GET", f"/admin/events?period={PERIOD}", token=tok)
check("event_log có DOC tiep_nhan lĩnh vực tthc/tai_chinh và ban_hanh y_te", {(e["event_type"], e["domain"]) for e in ev} >= {("tiep_nhan", "tthc"), ("tiep_nhan", "tai_chinh"), ("ban_hanh", "y_te")}, str([(e["object_id"], e["event_type"], e["domain"]) for e in ev]))
st, tasks = call("GET", "/admin/tasks?status=done", token=tok)
check("việc từ QLVB đã hoàn thành có source_type qlvb", any(t.get("source_type") == "qlvb" for t in tasks.get("data", [])), str([(t["code"], t.get("source_type"), t["status"]) for t in tasks.get("data", [])][:5]))
st, sub = call("POST", f"/admin/qlvb/report-runs/{rid}/submit", {"signer": "Chủ tịch UBND"}, token=tok)
check("trình ký: run submitted, có qlvb_doc_id", st in (200, 201) and sub["status"] == "submitted" and sub["qlvb_doc_id"], str(sub)[:300])
st, e = call("POST", f"/admin/report-runs/{rid}/unlock", {"reason": "x"}, token=tok)
check("kỳ đã trình ký không mở lại được", st == 400, str(e)[:160])
st, chk = call("POST", f"/admin/qlvb/report-runs/{rid}/check-sign", token=tok)
check("chưa ký → pending", chk["status"] == "pending", str(chk)[:200])
st, docs = call("GET", f"/admin/qlvb/documents?direction=di", token=tok)
pend = next(d for d in docs if d["sign_status"] == "pending" and d["run_id"] == rid)
st, sg = call("POST", f"/admin/qlvb/sim/sign/{pend['id']}", {"nguoi_ky": "Chủ tịch UBND"}, token=tok)
check("giả lập ký: cấp số ký hiệu /BC-UBND", st in (200, 201) and sg["sign_status"] == "signed" and sg["so_ky_hieu"].endswith("/BC-UBND"), str(sg)[:300])
st, chk = call("POST", f"/admin/qlvb/report-runs/{rid}/check-sign", token=tok)
check("nhận bản ký: run signed, có số/ngày/file", chk["status"] == "signed" and chk["run"]["status"] == "signed" and chk["run"]["signed_number"] == sg["so_ky_hieu"] and chk["run"]["signed_file"], str(chk)[:300])
st, ev = call("GET", f"/admin/events?period=2026-M{dt.date.today().month:02d}", token=tok)
check("sự kiện ban hành báo cáo đã ký ghi vào event_log", any(e["object_id"] == sg["so_ky_hieu"] and e["event_type"] == "ban_hanh" for e in ev), str([(e["object_id"], e["event_type"]) for e in ev][:6]))
st, e = call("POST", "/admin/facts/ingest", {"records": [{"indicator_id": "tthc.tiep_nhan", "period": PERIOD, "as_of_date": CUTOFF, "value": 1, "capture_method": "form"}]}, token=tok)
check("kỳ đã ký: số vẫn locked, không nhận số mới", st == 400, str(e)[:160])

print("4. KPI (D6)")
st, ce = call("POST", "/admin/kpi/employees/sync-from-accounts", token=tok)
check("đồng bộ cán bộ từ tài khoản", st in (200, 201) and ce["created"] + ce["updated"] >= 1, str(ce))
st, crit = call("GET", "/admin/kpi/criteria", token=tok)
check("4 tiêu chí từ config/kpi/criteria.yaml", st == 200 and len(crit) == 4 and any(c["criterion_id"] == "kpi.viec_dung_han" for c in crit), str([c["criterion_id"] for c in crit]))
st, ex = call("GET", f"/admin/kpi/export/tasks?period={PERIOD}", token=tok)
check("xuất việc theo cán bộ cho module KPI (có việc từ QLVB)", st == 200 and any(r["source"] == "qlvb" for r in ex), str(ex)[:300])
st, cp = call("POST", "/admin/kpi/compute", {"period": PERIOD}, token=tok)
check("tính điểm đề xuất", st in (200, 201) and cp["written"] >= 2, str(cp))
st, sc = call("GET", f"/admin/kpi/scores?period={PERIOD}", token=tok)
mine = [s for s in sc if s["employee_code"] == EMAIL.lower()]
vd = next((s for s in mine if s["criterion_id"] == "kpi.van_ban_dung_han"), None)
check("văn bản đúng hạn: 1/2 văn bản xong đúng hạn → 50% × 30 = 15 điểm", vd and abs(vd["value"] - 50) < 0.01 and abs(vd["score"] - 15) < 0.01, str(vd)[:300])
kl = next((s for s in mine if s["criterion_id"] == "kpi.ky_luat_bao_cao"), None)
check("kỷ luật báo cáo: 7 số nộp (5 bóc + 1 sửa + 1 bị loại... ) có bằng chứng fact_ids", kl and kl["value"] is not None and kl["evidence_ref"].get("fact_ids"), str(kl)[:300])
check("điểm ở trạng thái proposed", all(s["status"] == "proposed" for s in mine), str([s["status"] for s in mine]))
st, exr = call("POST", f"/admin/kpi/scores/{vd['id']}/explain", {"explanation": "Văn bản 1250 còn trong hạn"}, token=tok)
check("giải trình → explained", st in (200, 201) and exr["status"] == "explained", str(exr)[:200])
st, ap = call("POST", "/admin/kpi/scores/approve", {"ids": [s["id"] for s in mine]}, token=tok)
check("lãnh đạo duyệt", st in (200, 201) and ap["approved"] == len(mine), str(ap))
st, cp2 = call("POST", "/admin/kpi/compute", {"period": PERIOD}, token=tok)
check("tính lại không ghi đè điểm đã duyệt", cp2["kept_approved"] >= len(mine), str(cp2))
st, im = call("POST", "/admin/kpi/scores/import", {"rows": [{"employee_code": EMAIL.lower(), "period": PERIOD, "criterion_id": "kpi.danh_gia_lanh_dao", "value": 90, "evidence_ref": {"note": "lãnh đạo nhập"}}]}, token=tok)
check("nhận điểm định tính (module KPI / lãnh đạo nhập)", st in (200, 201) and im["upserted"] == 1, str(im))
st, sm = call("GET", f"/admin/kpi/summary?period={PERIOD}", token=tok)
check("tổng hợp: có cán bộ được đánh giá, điểm trung bình", st == 200 and sm["totals"]["so_can_bo_danh_gia"] >= 1 and sm["totals"]["diem_trung_binh"] is not None, str(sm["totals"]))
st, pb = call("POST", "/admin/kpi/publish", {"period": PERIOD}, token=tok)
check("đẩy số KPI vào kho (extracted, capture api)", st in (200, 201) and len(pb) >= 2 and all(o["status"] == "extracted" for o in pb), str(pb)[:300])
st, au = call("GET", "/admin/audit-logs", token=tok)
check("lượt xem điểm KPI được ghi audit (module kpi, view_scores)", any(a.get("module") == "kpi" and a.get("action") == "view_scores" for a in (au if isinstance(au, list) else au.get("data", []))), "")

print("5. Dashboard lãnh đạo")
st, ov = call("GET", f"/admin/dashboard/overview?period={PERIOD}", token=tok)
check("overview: run signed, lĩnh vực tthc đủ 5 số duyệt, văn bản 2 đến 1 đi, KPI có", st == 200 and ov["run"]["status"] == "signed" and any(d["domain"] == "tthc" and d["approved"] == 5 for d in ov["domains"]) and ov["documents"]["den"] == 2 and ov["documents"]["di"] >= 1 and ov["kpi"]["so_can_bo_danh_gia"] >= 1, str(ov)[:500])
st, tr = call("GET", f"/admin/dashboard/trends?indicators=tthc.tiep_nhan,tthc.giai_quyet&period={PERIOD}&n=6", token=tok)
check("trends: 6 kỳ tháng, giá trị tháng 8 = 5.000 (locked)", st == 200 and len(tr["periods"]) == 6 and next(p for s in tr["series"] if s["indicator_id"] == "tthc.tiep_nhan" for p in s["points"] if p["period"] == PERIOD)["value"] == 5000, str(tr)[:300])
st, pk = call("GET", f"/admin/dashboard/project-kpis?period={PERIOD}", token=tok)
check("5 KPI dự án: bậc 1–3 = 100% (5 export + 3 api), slot tự điền 100%, 1/8 số tự động phải sửa → 87.5%, 0 lỗi lọt bản ký, có giờ tổng hợp",
      st == 200 and pk["kpis"]["ty_le_bac_1_3"] == 100 and pk["kpis"]["slot_tu_dien"] == 100 and pk["kpis"]["so_boc_khong_can_sua"] == 87.5 and pk["kpis"]["loi_doi_soat_lot_ban_ky"] == 0 and pk["kpis"]["thoi_gian_tong_hop_gio"] is not None, str(pk))

print(); print("KẾT QUẢ:", "ĐẠT" if not FAILS else f"CÒN {len(FAILS)} LỖI: {FAILS}")
sys.exit(1 if FAILS else 0)
