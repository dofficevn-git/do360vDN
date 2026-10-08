# -*- coding: utf-8 -*-
"""Biểu mẫu Excel chuẩn (tools/make_forms.py): điền số vào mẫu sinh ra → luật sinh kèm phải bóc đúng 100%, không cần server.
   python eval/test_forms.py"""
import glob, io, os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "services", "reader"))
import yaml
from openpyxl import load_workbook
from app.readers import read_any
from app.rules import apply_rule

ROOT = os.path.join(os.path.dirname(__file__), "..")
FAILS = []
def check(name, cond, detail=""):
    print(("  OK   " if cond else "  FAIL ") + name + (f" — {detail}" if detail and not cond else ""))
    if not cond: FAILS.append(name)

forms = sorted(glob.glob(os.path.join(ROOT, "config", "templates", "forms", "*.xlsx")))
check("có ít nhất một biểu mẫu sinh ra", len(forms) > 0, "chạy python tools/make_forms.py")
for f in forms:
    dept = os.path.basename(f)[len("bieu-mau-"):-5].replace("-", "_")
    rule_path = os.path.join(ROOT, "config", "extract-rules", f"bieu_mau_{dept}_xlsx.yaml")
    check(f"{dept}: có luật đi kèm", os.path.exists(rule_path))
    if not os.path.exists(rule_path): continue
    rule = yaml.safe_load(io.open(rule_path, encoding="utf-8"))
    wb = load_workbook(f); ws = wb["Biểu mẫu"]
    hdr = next(r for r in range(1, 12) if ws.cell(row=r, column=2).value == "Mã chỉ tiêu")
    expected, types = {}, {}
    for r in range(hdr + 1, ws.max_row + 1):
        ind = ws.cell(row=r, column=2).value
        if not ind: continue
        v = 1000 + r * 7 if (r % 3) else 12.5
        ws.cell(row=r, column=5).value = v; expected[ind] = float(v); types[ind] = "decimal"
    buf = io.BytesIO(); wb.save(buf)
    doc = read_any(os.path.basename(f), buf.getvalue(), allow_ocr=False)
    res = apply_rule(doc, rule, types)
    got = {e.indicator_id: float(e.value) for e in res.records}
    check(f"{dept}: luật khớp, bóc đủ {len(expected)} chỉ tiêu", res.matched and got == expected, f"got={got} misses={res.misses}")
    check(f"{dept}: mọi bản ghi có ô nguồn", all(e.source_ref.get("cell") for e in res.records))

print(); print("KẾT QUẢ:", "ĐẠT" if not FAILS else f"CÒN {len(FAILS)} LỖI: {FAILS}")
sys.exit(1 if FAILS else 0)
