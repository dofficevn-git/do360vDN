# -*- coding: utf-8 -*-
"""T2 — Đọc tài liệu: Excel, Word, PDF có lớp chữ đọc thẳng; ảnh/PDF scan qua OCR (nếu cài).
Mọi reader trả về cùng một cấu trúc `Doc` để bộ máy luật (rules.py) không cần biết nguồn."""
from __future__ import annotations
import io
from dataclasses import dataclass, field
from typing import Optional


@dataclass
class Cell:
    row: int        # 1-based
    col: int        # 1-based
    text: str
    ref: str        # "C5"


@dataclass
class Table:
    name: str                       # tên sheet hoặc "table-1"
    cells: list[Cell] = field(default_factory=list)
    page: Optional[int] = None

    def grid(self) -> dict[tuple[int, int], Cell]:
        return {(c.row, c.col): c for c in self.cells}


@dataclass
class TextBlock:
    text: str
    page: Optional[int] = None
    bbox: Optional[tuple[float, float, float, float]] = None
    confidence: float = 1.0


@dataclass
class Doc:
    file_kind: str                  # xlsx | docx | pdf_text | pdf_scan | image
    tables: list[Table] = field(default_factory=list)
    blocks: list[TextBlock] = field(default_factory=list)
    pages: int = 1
    ocr: bool = False
    warnings: list[str] = field(default_factory=list)

    def full_text(self) -> str:
        return "\n".join(b.text for b in self.blocks)


def detect_kind(name: str, data: bytes) -> str:
    n = (name or "").lower()
    if n.endswith((".xlsx", ".xlsm")) or data[:4] == b"PK\x03\x04" and n.endswith(".xlsx"):
        return "xlsx"
    if n.endswith(".docx"):
        return "docx"
    if n.endswith(".pdf") or data[:5] == b"%PDF-":
        return "pdf"
    if n.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")):
        return "image"
    if n.endswith(".csv"):
        return "csv"
    return "other"


def read_xlsx(data: bytes) -> Doc:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    doc = Doc(file_kind="xlsx")
    for ws in wb.worksheets:
        t = Table(name=ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                t.cells.append(Cell(c.row, c.column, str(c.value).strip(), c.coordinate))
        doc.tables.append(t)
    doc.pages = len(wb.worksheets)
    return doc


def read_csv(data: bytes) -> Doc:
    import csv
    text = data.decode("utf-8-sig", errors="replace")
    t = Table(name="csv")
    for r, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        for cidx, v in enumerate(row, start=1):
            if v.strip():
                t.cells.append(Cell(r, cidx, v.strip(), f"R{r}C{cidx}"))
    return Doc(file_kind="xlsx", tables=[t])


def read_docx(data: bytes) -> Doc:
    from docx import Document
    d = Document(io.BytesIO(data))
    doc = Doc(file_kind="docx")
    for p in d.paragraphs:
        if p.text.strip():
            doc.blocks.append(TextBlock(p.text.strip()))
    for ti, tbl in enumerate(d.tables, start=1):
        t = Table(name=f"table-{ti}")
        for r, row in enumerate(tbl.rows, start=1):
            for c, cell in enumerate(row.cells, start=1):
                if cell.text.strip():
                    t.cells.append(Cell(r, c, cell.text.strip(), f"R{r}C{c}"))
        doc.tables.append(t)
    return doc


def read_pdf(data: bytes, allow_ocr: bool = True) -> Doc:
    import pymupdf
    pdf = pymupdf.open(stream=data, filetype="pdf")
    doc = Doc(file_kind="pdf_text", pages=len(pdf))
    text_chars = 0
    for pno, page in enumerate(pdf, start=1):
        for b in page.get_text("blocks"):
            x0, y0, x1, y1, txt = b[0], b[1], b[2], b[3], b[4]
            txt = " ".join(str(txt).split())
            if txt:
                text_chars += len(txt)
                doc.blocks.append(TextBlock(txt, page=pno, bbox=(x0, y0, x1, y1)))
        try:
            for ti, tb in enumerate(page.find_tables().tables, start=1):
                t = Table(name=f"p{pno}-table-{ti}", page=pno)
                for r, row in enumerate(tb.extract(), start=1):
                    for c, v in enumerate(row, start=1):
                        if v and str(v).strip():
                            t.cells.append(Cell(r, c, " ".join(str(v).split()), f"R{r}C{c}"))
                if t.cells:
                    doc.tables.append(t)
        except Exception as err:  # find_tables không có ở bản cũ
            doc.warnings.append(f"trang {pno}: không nhận dạng bảng ({err})")
    # Lớp chữ quá mỏng so với số trang → bản scan
    if text_chars < 40 * len(pdf):
        if not allow_ocr:
            doc.file_kind = "pdf_scan"
            doc.warnings.append("PDF không có lớp chữ — cần OCR")
            return doc
        images = [page.get_pixmap(dpi=200).tobytes("png") for page in pdf]
        ocr_doc = read_images(images)
        ocr_doc.file_kind = "pdf_scan"
        return ocr_doc
    return doc


def ocr_available() -> bool:
    try:
        import paddleocr  # noqa: F401
        return True
    except Exception:
        return False


_OCR = None


def _ocr():
    global _OCR
    if _OCR is None:
        from paddleocr import PaddleOCR
        _OCR = PaddleOCR(use_angle_cls=True, lang="vi", show_log=False)
    return _OCR


def read_images(images: list[bytes]) -> Doc:
    """OCR từng ảnh → khối chữ có toạ độ + độ tin cậy. Không có PaddleOCR → ném lỗi rõ."""
    if not ocr_available():
        raise RuntimeError("OCR chưa cài trên máy này (paddleocr). Nộp Excel/Word/PDF có lớp chữ, hoặc cài theo services/reader/README.md.")
    import numpy as np
    from PIL import Image
    doc = Doc(file_kind="image", pages=len(images), ocr=True)
    for pno, img in enumerate(images, start=1):
        arr = np.array(Image.open(io.BytesIO(img)).convert("RGB"))
        for line in (_ocr().ocr(arr, cls=True) or [[]])[0] or []:
            box, (txt, conf) = line
            xs = [p[0] for p in box]; ys = [p[1] for p in box]
            doc.blocks.append(TextBlock(" ".join(str(txt).split()), page=pno, bbox=(min(xs), min(ys), max(xs), max(ys)), confidence=float(conf)))
    return doc


def read_any(name: str, data: bytes, allow_ocr: bool = True) -> Doc:
    kind = detect_kind(name, data)
    if kind == "xlsx":
        return read_xlsx(data)
    if kind == "csv":
        return read_csv(data)
    if kind == "docx":
        return read_docx(data)
    if kind == "pdf":
        return read_pdf(data, allow_ocr)
    if kind == "image":
        if not allow_ocr:
            raise RuntimeError("ảnh cần OCR")
        return read_images([data])
    raise RuntimeError(f"không đọc được loại file: {name}")
